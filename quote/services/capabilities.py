"""
Supplier capabilities -- which NSNs / whole FSCs each supplier can supply.

These lists feed solicitation matching (``services/matching``). This module is
everything a rep can *do* to them:

  * read a pasted list or an uploaded CSV / Excel file into a table,
  * work out which columns hold NSNs / FSCs and which hold the supplier,
  * resolve supplier names / CAGEs against the directory,
  * build an ``ImportPlan`` -- a dry run that says what is new, what is already
    on file, what could not be read, and what it would do to open solicitations,
  * commit a plan (recording a ``QuoteCapabilityImport`` so it can be undone),
  * remove capabilities and undo an import, pruning the matches they created.

Views only orchestrate; every rule lives here. ``build_plan`` never writes.

SQL Server notes: ``__in`` lists are chunked under the 2,100-parameter limit,
inserts are chunked at 200 (what the importer uses in production) and nothing
iterates a lazy queryset while writing.
"""
import csv
import io
import logging
import re
from dataclasses import dataclass, field

from django.db import IntegrityError, transaction
from django.db.models import Count, Max
from django.utils import timezone

from dibbs.models import SolicitationLine
from products.nsn_utils import format_nsn
from quote.models import (
    QuoteCapabilityImport,
    QuoteSolicitation,
    QuoteSolicitationMatch,
    QuoteSupplierFSC,
    QuoteSupplierNSN,
)
from quote.services.matching import (
    IN_CHUNK,
    INSERT_CHUNK,
    _chunked,
    match_solicitations,
    normalize_nsn,
    open_matchable_ids,
    preview_matches,
    prune_derived_matches,
)
from suppliers.models import Supplier, SupplierAlias

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
DETECT_ROWS = 2000        # rows sampled when guessing columns
SAMPLE_ROWS = 6           # rows echoed back for the mapping strip
UNREADABLE_CAP = 200      # unreadable items echoed back to the browser
SUPPLIER_LIST_CAP = 400   # supplier rows echoed back from a preview
NSN_SAMPLE = 5

ROLE_IGNORE = 'ignore'
ROLE_CAGE = 'cage'
ROLE_NAME = 'name'
ROLE_ITEMS = 'items'
ROLES = (ROLE_IGNORE, ROLE_CAGE, ROLE_NAME, ROLE_ITEMS)

SKIP = 'skip'             # an assignment value: "leave this supplier out"


class CapabilityError(Exception):
    """A problem worth showing the rep verbatim (bad file, nothing to import...)."""


# ── Finding NSNs and FSCs in text ────────────────────────────────────────────

_DASH_CHARS = dict.fromkeys(map(ord, '‐‑‒–—―−'), '-')

# 13 digits, optionally grouped 4-2-3-4 with hyphens or spaces, not part of a longer number.
_NSN_RE = re.compile(r'(?<!\d)(\d{4})[- ]?(\d{2})[- ]?(\d{3})[- ]?(\d{4})(?!\d)')
# A bare four-digit token (not the front of 5340-01-...).
_FSC_RE = re.compile(r'(?<!\d)(?<!\d-)(\d{4})(?!\d)(?!-\d)')
# What is left of a number-ish run once NSNs and FSCs are removed.
_RUN_RE = re.compile(r'(?<!\d)\d[\d-]{6,}\d(?!\d)')
# 5.34E+12 -- Excel's scientific notation, digits already lost.
_SCI_RE = re.compile(r'(?<![\w.])\d(?:\.\d+)?[eE]\+?\d{1,2}(?![\w])')


def find_items(text):
    """
    Pull NSNs and FSCs out of a piece of text.

    Returns ``(nsns, fscs, unreadable)`` -- 13-digit NSNs, 4-digit FSCs, and
    ``[(text, reason)]`` for number-shaped things that are neither (a NIIN with
    no FSC, an NSN a digit short, Excel scientific notation...). Words are ignored.
    """
    cleaned = (text or '').translate(_DASH_CHARS).replace(' ', ' ')
    unreadable = []

    nsns = [''.join(m.groups()) for m in _NSN_RE.finditer(cleaned)]
    rest = _NSN_RE.sub(' ', cleaned)

    for m in _SCI_RE.finditer(rest):
        unreadable.append((
            m.group(0),
            'Excel turned this number into scientific notation. Re-export the '
            'column formatted as Text.',
        ))
    rest = _SCI_RE.sub(' ', rest)

    fscs = []
    for m in _FSC_RE.finditer(rest):
        fsc = m.group(1)
        if fsc[:2] >= '10':
            fscs.append(fsc)
        else:
            unreadable.append((fsc, 'Not an FSC (classes start at 10xx).'))
    rest = _FSC_RE.sub(' ', rest)

    for m in _RUN_RE.finditer(rest):
        run = m.group(0)
        digit_string = re.sub(r'\D', '', run)
        digits = len(digit_string)
        if digits < 9:              # dates, short codes: not worth flagging
            continue
        if digits == 13:            # an NSN with odd hyphenation
            nsns.append(digit_string)
            continue
        if digits == 9:
            reason = 'NIIN only (9 digits). It needs the 4-digit FSC in front to be an NSN.'
        else:
            reason = f'{digits} digits. An NSN has 13.'
        unreadable.append((run, reason))
    return nsns, fscs, unreadable


# ── Reading a paste or an upload into a table ────────────────────────────────

@dataclass
class Table:
    rows: list
    source_name: str = ''
    note: str = ''
    truncated: bool = False

    @property
    def width(self):
        return max((len(r) for r in self.rows), default=0)


def _sniff_delimiter(lines):
    """Tab, semicolon, comma or pipe when most of the sample lines use it."""
    if not lines:
        return None
    best, best_lines = None, 0
    for delim in ('\t', ';', ',', '|'):
        n = sum(1 for line in lines if delim in line)
        if n > best_lines:
            best, best_lines = delim, n
    return best if best_lines * 2 >= len(lines) else None


def read_text(text, source_name='Pasted text'):
    """A paste or decoded text file as a table (one column when nothing delimits it)."""
    text = (text or '').replace('\r\n', '\n').replace('\r', '\n')
    lines = [line for line in text.split('\n') if line.strip()]
    delim = _sniff_delimiter(lines[:20])
    if delim:
        parsed = csv.reader(io.StringIO('\n'.join(lines)), delimiter=delim)
    else:
        parsed = ([line] for line in lines)
    rows, truncated = [], False
    for cells in parsed:
        cells = [c.strip() for c in cells]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            continue
        rows.append(cells)
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
    return Table(rows=rows, source_name=source_name, truncated=truncated)


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f'{value:.10g}'
    return str(value).strip()


def _read_xlsx(data, source_name):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError...
        raise CapabilityError("Couldn't open that Excel file. Is it a real .xlsx?") from exc
    try:
        sheets = list(workbook.worksheets)
        for sheet in sheets:
            try:
                sheet.reset_dimensions()   # some writers leave a bogus <dimension> tag
            except Exception:
                pass
            rows, truncated = [], False
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not any(cells):
                    continue
                rows.append(cells)
                if len(rows) >= MAX_ROWS:
                    truncated = True
                    break
            if rows:
                others = len(sheets) - 1
                note = f'Read sheet "{sheet.title}"'
                if others:
                    note += f' ({others} other sheet{"s" if others != 1 else ""} ignored)'
                return Table(rows=rows, source_name=source_name, note=note, truncated=truncated)
    finally:
        workbook.close()
    return Table(rows=[], source_name=source_name)


def _decode(data):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode('latin-1')


def read_upload(uploaded):
    """An uploaded CSV / TSV / TXT / XLSX file as a table."""
    name = getattr(uploaded, 'name', '') or 'upload'
    if getattr(uploaded, 'size', 0) > MAX_UPLOAD_BYTES:
        raise CapabilityError('That file is bigger than 10 MB. Split it up and import in parts.')
    ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    data = uploaded.read()
    if ext in ('xlsx', 'xlsm'):
        return _read_xlsx(data, name)
    if ext == 'xls':
        raise CapabilityError('Old .xls files are not supported. Save it as .xlsx or CSV and try again.')
    if ext in ('csv', 'tsv', 'txt', ''):
        return read_text(_decode(data), source_name=name)
    raise CapabilityError(f'.{ext} files are not supported. Use CSV, TXT or Excel (.xlsx).')


# ── Suppliers: matching names and CAGEs to the directory ─────────────────────

_NAME_SUFFIXES = {
    'INC', 'INCORPORATED', 'LLC', 'LTD', 'LIMITED', 'CORP', 'CORPORATION',
    'CO', 'COMPANY', 'LP', 'LLP', 'PLLC',
}


def norm_name(value):
    """Upper-case, punctuation-free, corporate suffix dropped: 'Acme Tool, Inc.' -> 'ACME TOOL'."""
    text = re.sub(r'[^A-Z0-9]+', ' ', (value or '').upper().replace('&', ' AND '))
    tokens = text.split()
    while tokens and tokens[-1] in _NAME_SUFFIXES:
        tokens.pop()
    if tokens and tokens[0] == 'THE':
        tokens.pop(0)
    return ' '.join(tokens)


@dataclass
class Resolution:
    status: str                  # ok | ambiguous | archived | unknown | blank
    supplier_id: int = None
    candidates: list = field(default_factory=list)


class SupplierResolver:
    """
    Maps the supplier text in an import row (a CAGE, a name, or both) to a
    directory supplier. One load of the directory, then dictionary lookups.
    CAGE wins over name; a name must match exactly one active supplier.
    """

    def __init__(self):
        self.info = {}
        self.by_cage = {}
        self.by_name = {}
        for row in Supplier.objects.values('id', 'name', 'cage_code', 'archived'):
            self.info[row['id']] = row
            cage = (row['cage_code'] or '').strip().upper()
            if cage:
                self.by_cage[cage] = row['id']
            key = norm_name(row['name'])
            if key:
                self.by_name.setdefault(key, set()).add(row['id'])
        for row in SupplierAlias.objects.values('supplier_id', 'name'):
            key = norm_name(row['name'])
            if key and row['supplier_id'] in self.info:
                self.by_name.setdefault(key, set()).add(row['supplier_id'])

    def has_cage(self, value):
        return (value or '').strip().upper() in self.by_cage

    def _candidate(self, sid):
        row = self.info[sid]
        return {'id': sid, 'name': row['name'] or '(no name)', 'cage': row['cage_code'] or ''}

    def resolve(self, cage='', name=''):
        cage = (cage or '').strip().upper()
        name_key = norm_name(name)
        if not cage and not name_key:
            return Resolution('blank')

        if cage and cage in self.by_cage:
            sid = self.by_cage[cage]
            if self.info[sid]['archived']:
                return Resolution('archived', candidates=[self._candidate(sid)])
            return Resolution('ok', sid)

        ids = sorted(self.by_name.get(name_key, ())) if name_key else []
        active = [i for i in ids if not self.info[i]['archived']]
        if len(active) == 1:
            return Resolution('ok', active[0])
        if len(active) > 1:
            return Resolution('ambiguous', candidates=[self._candidate(i) for i in active[:6]])
        if ids:
            return Resolution('archived', candidates=[self._candidate(i) for i in ids[:6]])
        return Resolution('unknown')


# ── Column mapping ───────────────────────────────────────────────────────────

@dataclass
class Mapping:
    header: bool
    roles: list

    def as_dict(self):
        return {'header': self.header, 'roles': list(self.roles)}

    @classmethod
    def from_payload(cls, payload, width):
        """A browser-supplied mapping, validated; ``None`` when unusable."""
        if not isinstance(payload, dict):
            return None
        roles = payload.get('roles')
        if not isinstance(roles, list):
            return None
        clean, seen = [], set()
        for j in range(width):
            role = roles[j] if j < len(roles) else ROLE_IGNORE
            if role not in ROLES:
                role = ROLE_IGNORE
            if role in (ROLE_CAGE, ROLE_NAME):
                if role in seen:
                    role = ROLE_IGNORE
                seen.add(role)
            clean.append(role)
        return cls(header=bool(payload.get('header')), roles=clean)

    def cols(self, role):
        return [j for j, r in enumerate(self.roles) if r == role]


_CAGE_SHAPE = re.compile(r'^[A-Z0-9]{5}$')
_CAGE_HEADER = re.compile(r'\bcage\b', re.I)
_NAME_HEADER = re.compile(r'supplier|vendor|company|manufacturer|mfr|firm|source|name', re.I)


def _has_items(cell):
    nsns, fscs, _ = find_items(cell)
    return bool(nsns or fscs)


def detect_mapping(table, resolver, single_supplier):
    """
    Guess the header row and each column's role from the data itself -- a column
    of NSN/FSC-looking cells is items; one whose cells are CAGEs (or names) we
    know is the supplier -- with header text as a tie-breaker.
    """
    rows, width = table.rows, table.width
    header = bool(
        rows
        and any(c.strip() for c in rows[0])
        and not any(_has_items(c) for c in rows[0])
        and any(_has_items(c) for r in rows[1:6] for c in r)
    )
    body = rows[1:1 + DETECT_ROWS] if header else rows[:DETECT_ROWS]
    head = rows[0] if header else []

    def column(j):
        return [r[j].strip() for r in body if j < len(r) and r[j].strip()]

    roles = [ROLE_IGNORE] * width
    items = set()
    for j in range(width):
        cells = column(j)
        if cells and sum(_has_items(c) for c in cells) / len(cells) >= 0.5:
            items.add(j)
            roles[j] = ROLE_ITEMS

    if not single_supplier:
        best = {ROLE_CAGE: (0.0, None), ROLE_NAME: (0.0, None)}
        for j in range(width):
            if j in items:
                continue
            cells = column(j)
            if not cells:
                continue
            cage_score = sum(resolver.has_cage(c) for c in cells) / len(cells)
            name_score = sum(
                1 for c in cells if resolver.resolve(name=c).status == 'ok'
            ) / len(cells)
            if cage_score > best[ROLE_CAGE][0]:
                best[ROLE_CAGE] = (cage_score, j)
            if name_score > best[ROLE_NAME][0]:
                best[ROLE_NAME] = (name_score, j)
        cage_col = best[ROLE_CAGE][1] if best[ROLE_CAGE][0] >= 0.3 else None
        name_col = best[ROLE_NAME][1] if best[ROLE_NAME][0] >= 0.3 else None
        if name_col == cage_col:
            name_col = None

        # Nothing in the directory recognised? Fall back to header text / CAGE shape.
        for j in range(width):
            if j in items or j in (cage_col, name_col):
                continue
            label = head[j] if j < len(head) else ''
            cells = column(j)
            if cage_col is None and label and _CAGE_HEADER.search(label):
                cage_col = j
            elif cage_col is None and cells and sum(
                bool(_CAGE_SHAPE.match(c.upper())) for c in cells
            ) / len(cells) >= 0.7:
                cage_col = j
            elif name_col is None and label and _NAME_HEADER.search(label):
                name_col = j
        if cage_col is not None:
            roles[cage_col] = ROLE_CAGE
        if name_col is not None and name_col != cage_col:
            roles[name_col] = ROLE_NAME

    if ROLE_ITEMS not in roles:
        # No column looked like NSNs/FSCs: read every column that isn't the supplier.
        roles = [r if r in (ROLE_CAGE, ROLE_NAME) else ROLE_ITEMS for r in roles]
    return Mapping(header=header, roles=roles)


def _column_labels(table, mapping):
    head = table.rows[0] if mapping.header and table.rows else []
    labels = []
    for j in range(table.width):
        text = head[j].strip() if j < len(head) else ''
        letter = chr(ord('A') + j) if j < 26 else str(j + 1)
        labels.append(text or f'Column {letter}')
    return labels


# ── Plan (dry run) ───────────────────────────────────────────────────────────

@dataclass
class SupplierPlan:
    supplier: dict
    nsns: set = field(default_factory=set)
    fscs: set = field(default_factory=set)
    new_nsns: list = field(default_factory=list)
    new_fscs: list = field(default_factory=list)
    existing_nsns: int = 0
    existing_fscs: int = 0


@dataclass
class ImportPlan:
    mode: str
    source_name: str
    note: str = ''
    truncated: bool = False
    table: Table = None
    mapping: Mapping = None
    rows_read: int = 0
    rows_used: int = 0
    suppliers: dict = field(default_factory=dict)        # supplier id -> SupplierPlan
    unresolved: dict = field(default_factory=dict)       # key -> bucket
    unreadable: list = field(default_factory=list)
    unreadable_count: int = 0
    problems: list = field(default_factory=list)
    impact: dict = field(default_factory=dict)
    fsc_open: dict = field(default_factory=dict)

    @property
    def nsns_new(self):
        return sum(len(p.new_nsns) for p in self.suppliers.values())

    @property
    def fscs_new(self):
        return sum(len(p.new_fscs) for p in self.suppliers.values())

    @property
    def nsns_existing(self):
        return sum(p.existing_nsns for p in self.suppliers.values())

    @property
    def fscs_existing(self):
        return sum(p.existing_fscs for p in self.suppliers.values())

    @property
    def can_commit(self):
        return not self.problems and (self.nsns_new + self.fscs_new) > 0


def _open_solicitation_counts_by_fsc(fscs):
    """``{fsc: number of open, still-matchable solicitations with a line in it}``."""
    if not fscs:
        return {}
    matchable = QuoteSolicitation.objects.filter(
        status__in=QuoteSolicitation.MATCHING_STATES,
        solicitation__return_by_date__gte=timezone.now().date(),
    ).values('solicitation_id')
    counts = {}
    for chunk in _chunked(sorted(fscs), IN_CHUNK):
        rows = (
            SolicitationLine.objects.filter(solicitation_id__in=matchable, fsc__in=chunk)
            .values('fsc').annotate(n=Count('solicitation_id', distinct=True))
        )
        for row in rows:
            counts[row['fsc']] = row['n']
    return counts


def _existing_capabilities(supplier_ids):
    """``({sid: {nsn}}, {sid: {fsc}})`` already on file for these suppliers."""
    nsns, fscs = {}, {}
    for chunk in _chunked(sorted(supplier_ids), IN_CHUNK):
        for sid, nsn in QuoteSupplierNSN.objects.filter(supplier_id__in=chunk).values_list(
            'supplier_id', 'nsn'
        ):
            nsns.setdefault(sid, set()).add(nsn)
        for sid, fsc in QuoteSupplierFSC.objects.filter(supplier_id__in=chunk).values_list(
            'supplier_id', 'fsc'
        ):
            fscs.setdefault(sid, set()).add(fsc)
    return nsns, fscs


def build_plan(table, *, supplier=None, mapping=None, assignments=None, resolver=None,
               with_impact=True):
    """
    Dry run of an import. ``supplier`` fixes the supplier (one-supplier mode);
    otherwise each row's supplier comes from the mapped CAGE / name column.
    ``assignments`` is ``{key: supplier_id | 'skip'}`` for suppliers the
    directory could not resolve on its own. Nothing is written.
    """
    single = supplier is not None
    plan = ImportPlan(
        mode=QuoteCapabilityImport.MODE_SINGLE if single else QuoteCapabilityImport.MODE_MULTI,
        source_name=table.source_name, note=table.note, truncated=table.truncated, table=table,
    )
    if not table.rows:
        plan.problems.append("There's nothing in that to read.")
        return plan

    resolver = resolver or (None if single else SupplierResolver())
    if mapping is None or len(mapping.roles) != table.width:
        mapping = detect_mapping(table, resolver, single)
    plan.mapping = mapping

    item_cols = mapping.cols(ROLE_ITEMS)
    cage_cols = mapping.cols(ROLE_CAGE)
    name_cols = mapping.cols(ROLE_NAME)
    cage_col = cage_cols[0] if cage_cols else None
    name_col = name_cols[0] if name_cols else None
    if not item_cols:
        plan.problems.append('Pick the column(s) that hold NSNs or FSCs.')
    if not single and cage_col is None and name_col is None:
        if table.width < 2:
            plan.problems.append(
                'That is a plain list with no supplier named. Choose "One supplier" and pick '
                'which supplier it belongs to, or add a supplier column.'
            )
        else:
            plan.problems.append('Pick the column that holds the supplier (CAGE or name).')
    if plan.problems:
        return plan

    assignments = assignments or {}
    assigned_ids = {}
    wanted_ids = {int(v) for v in assignments.values() if str(v).isdigit()}
    if wanted_ids:
        assigned_ids = {
            row['id']: row for row in
            Supplier.objects.filter(pk__in=wanted_ids, archived=False).values('id', 'name', 'cage_code')
        }

    body = table.rows[1:] if mapping.header else table.rows
    per_supplier = {}
    resolutions = {}
    for offset, row in enumerate(body):
        plan.rows_read += 1
        nsns, fscs, junk = [], [], []
        for j in item_cols:
            if j < len(row):
                n, f, u = find_items(row[j])
                nsns += n
                fscs += f
                junk += u
        for text, reason in junk:
            plan.unreadable_count += 1
            if len(plan.unreadable) < UNREADABLE_CAP:
                plan.unreadable.append({
                    'row': offset + (2 if mapping.header else 1), 'text': text, 'reason': reason,
                })
        if not (nsns or fscs):
            continue

        if single:
            sid = supplier.pk
        else:
            cage = row[cage_col].strip() if cage_col is not None and cage_col < len(row) else ''
            name = row[name_col].strip() if name_col is not None and name_col < len(row) else ''
            key = f'{cage.upper()}|{norm_name(name)}'
            if key not in resolutions:
                resolutions[key] = resolver.resolve(cage, name)
            res = resolutions[key]
            if res.status == 'ok':
                sid = res.supplier_id
            else:
                bucket = plan.unresolved.get(key)
                if bucket is None:
                    label = ' / '.join(x for x in (name, cage) if x) or '(no supplier named)'
                    bucket = plan.unresolved[key] = {
                        'key': key, 'label': label, 'reason': res.status,
                        'candidates': res.candidates, 'rows': 0, 'items': 0,
                        'assigned': None, 'assigned_name': '',
                    }
                bucket['rows'] += 1
                bucket['items'] += len(nsns) + len(fscs)
                choice = assignments.get(key)
                if choice == SKIP:
                    bucket['assigned'] = SKIP
                    continue
                if str(choice).isdigit() and int(choice) in assigned_ids:
                    sid = int(choice)
                    bucket['assigned'] = sid
                    bucket['assigned_name'] = assigned_ids[sid]['name'] or ''
                else:
                    continue

        entry = per_supplier.setdefault(sid, (set(), set()))
        entry[0].update(nsns)
        entry[1].update(fscs)
        plan.rows_used += 1

    if not per_supplier and not plan.unresolved and not plan.unreadable_count:
        plan.problems.append("Couldn't find any NSNs or FSCs in that.")
        return plan

    infos = {}
    for chunk in _chunked(sorted(per_supplier), IN_CHUNK):
        for row in Supplier.objects.filter(pk__in=chunk).values('id', 'name', 'cage_code'):
            infos[row['id']] = row
    existing_nsns, existing_fscs = _existing_capabilities(per_supplier)
    for sid, (nsns, fscs) in per_supplier.items():
        info = infos.get(sid) or {'id': sid, 'name': '(unknown)', 'cage_code': ''}
        have_nsn = existing_nsns.get(sid, set())
        have_fsc = existing_fscs.get(sid, set())
        plan.suppliers[sid] = SupplierPlan(
            supplier=info, nsns=nsns, fscs=fscs,
            new_nsns=sorted(nsns - have_nsn), new_fscs=sorted(fscs - have_fsc),
            existing_nsns=len(nsns & have_nsn), existing_fscs=len(fscs & have_fsc),
        )

    if with_impact:
        _add_impact(plan)
    return plan


def _add_impact(plan):
    """What the new pairings would do to open solicitations (read-only)."""
    nsn_pairs = {(sid, n) for sid, sp in plan.suppliers.items() for n in sp.new_nsns}
    fsc_pairs = {(sid, f) for sid, sp in plan.suppliers.items() for f in sp.new_fscs}
    plan.fsc_open = _open_solicitation_counts_by_fsc({f for _, f in fsc_pairs})
    sol_ids = open_matchable_ids(
        nsns={n for _, n in nsn_pairs}, fscs={f for _, f in fsc_pairs},
    ) if (nsn_pairs or fsc_pairs) else set()
    plan.impact = (
        preview_matches(sorted(sol_ids), nsn_pairs, fsc_pairs) if sol_ids
        else {'new_links': 0, 'solicitations': 0, 'newly_matched': 0}
    )


def plan_to_preview(plan):
    """The plan as the JSON the preview screen renders."""
    table = plan.table
    mapping = plan.mapping
    out = {
        'mode': plan.mode,
        'source': plan.source_name,
        'note': plan.note,
        'truncated': plan.truncated,
        'problems': plan.problems,
        'can_commit': plan.can_commit,
        'summary': {
            'rows_read': plan.rows_read,
            'rows_used': plan.rows_used,
            'nsns_new': plan.nsns_new,
            'nsns_existing': plan.nsns_existing,
            'fscs_new': plan.fscs_new,
            'fscs_existing': plan.fscs_existing,
            'pairings_new': plan.nsns_new + plan.fscs_new,
            'unreadable': plan.unreadable_count,
            'suppliers': len(plan.suppliers),
            'unresolved': sum(1 for b in plan.unresolved.values() if b['assigned'] is None),
        },
        'impact': plan.impact,
        'unreadable': plan.unreadable,
    }
    if table is not None and mapping is not None:
        out['table'] = {
            'columns': _column_labels(table, mapping),
            'sample': [
                (r + [''] * (table.width - len(r)))[:table.width]
                for r in table.rows[1 if mapping.header else 0:][:SAMPLE_ROWS]
            ],
            'width': table.width,
            'rows': len(table.rows) - (1 if mapping.header else 0),
        }
        out['mapping'] = mapping.as_dict()

    suppliers = []
    ranked = sorted(
        plan.suppliers.values(),
        key=lambda p: (-(len(p.new_nsns) + len(p.new_fscs)), (p.supplier['name'] or '').lower()),
    )
    for sp in ranked[:SUPPLIER_LIST_CAP]:
        suppliers.append({
            'id': sp.supplier['id'],
            'name': sp.supplier['name'] or '(no name)',
            'cage': sp.supplier['cage_code'] or '',
            'nsns_new': len(sp.new_nsns),
            'nsns_existing': sp.existing_nsns,
            'fscs_new': len(sp.new_fscs),
            'fscs_existing': sp.existing_fscs,
            'sample': [format_nsn(n) for n in sp.new_nsns[:NSN_SAMPLE]],
            'fscs': [
                {'fsc': f, 'open': plan.fsc_open.get(f, 0)} for f in sp.new_fscs[:30]
            ],
        })
    out['suppliers'] = suppliers
    out['suppliers_hidden'] = max(0, len(plan.suppliers) - SUPPLIER_LIST_CAP)
    out['unresolved'] = list(plan.unresolved.values())
    return out


# ── Commit / remove / undo ───────────────────────────────────────────────────

def _bulk_insert(model, rows):
    """
    Insert ``rows`` in chunks. A chunk that trips the unique constraint (a
    teammate added the same pair a moment ago) falls back to one-by-one so the
    rest still lands. Returns how many were actually created.
    """
    created = 0
    for chunk in _chunked(rows, INSERT_CHUNK):
        try:
            with transaction.atomic():
                model.objects.bulk_create(chunk)
            created += len(chunk)
        except IntegrityError:
            for row in chunk:
                try:
                    with transaction.atomic():
                        row.save(force_insert=True)
                    created += 1
                except IntegrityError:
                    continue
    return created


def commit_plan(plan, user, *, supplier=None):
    """
    Write a plan: add the new pairings, record the import, then re-match open
    solicitations so the new capabilities show up in the queue straight away.
    """
    if plan.problems:
        raise CapabilityError(plan.problems[0])
    if not plan.can_commit:
        raise CapabilityError('Nothing new to add. Everything in that is already on file.')

    batch = QuoteCapabilityImport.objects.create(
        created_by=user, source_name=(plan.source_name or '')[:255], mode=plan.mode,
        supplier=supplier, suppliers_touched=len(plan.suppliers),
        nsns_existing=plan.nsns_existing, fscs_existing=plan.fscs_existing,
        unreadable=plan.unreadable_count,
    )
    nsn_rows = [
        QuoteSupplierNSN(supplier_id=sid, nsn=n, added_by=user, import_batch=batch)
        for sid, sp in plan.suppliers.items() for n in sp.new_nsns
    ]
    fsc_rows = [
        QuoteSupplierFSC(supplier_id=sid, fsc=f, added_by=user, import_batch=batch)
        for sid, sp in plan.suppliers.items() for f in sp.new_fscs
    ]
    batch.nsns_added = _bulk_insert(QuoteSupplierNSN, nsn_rows)
    batch.fscs_added = _bulk_insert(QuoteSupplierFSC, fsc_rows)

    summary = {}
    ids = open_matchable_ids(
        nsns={r.nsn for r in nsn_rows}, fscs={r.fsc for r in fsc_rows},
    )
    if ids:
        summary = match_solicitations(sorted(ids))
    batch.matches_created = summary.get('matches_created', 0)
    batch.solicitations_matched = summary.get('promoted_to_matched', 0)
    batch.save()
    logger.info(
        'capability import %s by %s: +%s NSN, +%s FSC, %s links, %s promoted',
        batch.pk, getattr(user, 'pk', None), batch.nsns_added, batch.fscs_added,
        batch.matches_created, batch.solicitations_matched,
    )
    return batch


def clean_keys(nsns=(), fscs=()):
    """Normalise browser-supplied NSNs / FSCs to stored form; junk is dropped."""
    nsn_set = {n for n in (normalize_nsn(x) for x in nsns) if n}
    fsc_set = {
        f for f in (''.join(ch for ch in str(x) if ch.isdigit()) for x in fscs)
        if len(f) == 4
    }
    return nsn_set, fsc_set


def remove_capabilities(supplier, nsns=(), fscs=()):
    """
    Remove NSNs / FSCs from a supplier, then prune the matches they justified
    on open, unworked solicitations. Returns counts for the toast.
    """
    nsn_set, fsc_set = clean_keys(nsns, fscs)
    removed_nsns, removed_fscs = set(), set()
    with transaction.atomic():
        for chunk in _chunked(sorted(nsn_set), IN_CHUNK):
            found = set(
                QuoteSupplierNSN.objects.filter(supplier=supplier, nsn__in=chunk)
                .values_list('nsn', flat=True)
            )
            if found:
                QuoteSupplierNSN.objects.filter(supplier=supplier, nsn__in=list(found)).delete()
                removed_nsns |= found
        for chunk in _chunked(sorted(fsc_set), IN_CHUNK):
            found = set(
                QuoteSupplierFSC.objects.filter(supplier=supplier, fsc__in=chunk)
                .values_list('fsc', flat=True)
            )
            if found:
                QuoteSupplierFSC.objects.filter(supplier=supplier, fsc__in=list(found)).delete()
                removed_fscs |= found
    pruned = {'matches_removed': 0, 'returned_to_unmatched': 0}
    if removed_nsns or removed_fscs:
        pruned = prune_derived_matches({supplier.pk: (removed_nsns, removed_fscs)})
    return {'nsns': len(removed_nsns), 'fscs': len(removed_fscs), **pruned}


def undo_import(batch, user):
    """Take back everything an import added (and only that), pruning its matches."""
    if batch.is_undone:
        raise CapabilityError('That import has already been undone.')
    removed = {}
    with transaction.atomic():
        nsn_rows = list(
            QuoteSupplierNSN.objects.filter(import_batch=batch).values_list('supplier_id', 'nsn')
        )
        fsc_rows = list(
            QuoteSupplierFSC.objects.filter(import_batch=batch).values_list('supplier_id', 'fsc')
        )
        for sid, nsn in nsn_rows:
            removed.setdefault(sid, (set(), set()))[0].add(nsn)
        for sid, fsc in fsc_rows:
            removed.setdefault(sid, (set(), set()))[1].add(fsc)
        QuoteSupplierNSN.objects.filter(import_batch=batch).delete()
        QuoteSupplierFSC.objects.filter(import_batch=batch).delete()
        batch.undone_at = timezone.now()
        batch.undone_by = user
        batch.save(update_fields=['undone_at', 'undone_by'])
    pruned = prune_derived_matches(removed) if removed else {
        'matches_removed': 0, 'returned_to_unmatched': 0,
    }
    return {'nsns': len(nsn_rows), 'fscs': len(fsc_rows), **pruned}


# ── Reads for the pages ──────────────────────────────────────────────────────

def capability_counts(supplier_ids=None):
    """``{supplier_id: {'nsns': n, 'fscs': n, 'last': datetime}}`` (suppliers with any)."""
    def grouped(model):
        qs = model.objects.all()
        if supplier_ids is not None:
            qs = qs.filter(supplier_id__in=list(supplier_ids))
        return {
            row['supplier_id']: row
            for row in qs.values('supplier_id').annotate(n=Count('id'), last=Max('added_at'))
        }

    nsns, fscs = grouped(QuoteSupplierNSN), grouped(QuoteSupplierFSC)
    out = {}
    for sid in set(nsns) | set(fscs):
        stamps = [r['last'] for r in (nsns.get(sid), fscs.get(sid)) if r and r['last']]
        out[sid] = {
            'nsns': nsns[sid]['n'] if sid in nsns else 0,
            'fscs': fscs[sid]['n'] if sid in fscs else 0,
            'last': max(stamps) if stamps else None,
        }
    return out


def capability_totals(counts=None):
    """``{'suppliers', 'nsns', 'fscs'}`` across everything on file."""
    counts = capability_counts() if counts is None else counts
    return {
        'suppliers': len(counts),
        'nsns': sum(c['nsns'] for c in counts.values()),
        'fscs': sum(c['fscs'] for c in counts.values()),
    }


def capability_overview():
    """Everything the Capabilities page lists: suppliers with capabilities + totals."""
    counts = capability_counts()
    rows = []
    for chunk in _chunked(sorted(counts), IN_CHUNK):
        for row in Supplier.objects.filter(pk__in=chunk).values('id', 'name', 'cage_code', 'archived'):
            rows.append({**row, **counts[row['id']]})
    rows.sort(key=lambda r: (r['name'] or '').lower())
    return rows, capability_totals(counts)


def linked_open_solicitations(supplier):
    """How many open, still-matchable solicitations this supplier is linked to."""
    return (
        QuoteSolicitationMatch.objects.filter(
            supplier=supplier,
            solicitation__return_by_date__gte=timezone.now().date(),
            solicitation__quote_state__status__in=list(QuoteSolicitation.MATCHING_STATES),
        ).values('solicitation_id').distinct().count()
    )


def export_rows(supplier=None):
    """
    Yield CSV rows (header first): one line per NSN / FSC pairing. Keyset-paged
    rather than ``.iterator()`` so it is safe on the MARS-less SQL Server setup.
    """
    yield ['Supplier', 'CAGE', 'Type', 'Code', 'Notes', 'Added', 'Added by', 'Import']
    suppliers = {}

    def label(sid):
        if sid not in suppliers:
            row = Supplier.objects.filter(pk=sid).values('name', 'cage_code').first() or {}
            suppliers[sid] = (row.get('name') or '', row.get('cage_code') or '')
        return suppliers[sid]

    for model, kind, attr in ((QuoteSupplierNSN, 'NSN', 'nsn'), (QuoteSupplierFSC, 'FSC', 'fsc')):
        last_pk = 0
        while True:
            qs = model.objects.filter(pk__gt=last_pk).select_related('added_by', 'import_batch')
            if supplier is not None:
                qs = qs.filter(supplier=supplier)
            page = list(qs.order_by('pk')[:2000])
            if not page:
                break
            for row in page:
                name, cage = label(row.supplier_id)
                code = getattr(row, attr)
                user = row.added_by
                yield [
                    name, cage, kind, format_nsn(code) if kind == 'NSN' else code,
                    row.notes, row.added_at.date().isoformat(),
                    (user.get_full_name() or user.get_username()) if user else '',
                    row.import_batch.source_name if row.import_batch else '',
                ]
            last_pk = page[-1].pk
