"""Supplier Research form, access, shell, panels, services, and Excel export."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from dibbs.models import DibbsAward
from dibbs.models.approved_sources import ApprovedSource
from quote.forms import CageLookupForm
from quote.services import supplier_research as sr
from users.models import AppPermission, AppRegistry

PANEL_KEYS = ('status', 'sam', 'awards', 'approved')

EMPTY_SUMMARY = {
    'total_awards': 0,
    'distinct_nsns': 0,
    'total_value': None,
    'first_award': None,
    'last_award': None,
    'data_current_through': None,
}
SAM_NOT_FOUND = {'state': 'not_found', 'rows': [], 'last_fetched': None}


def _quote_user(username='rep'):
    user = User.objects.create_user(username, password='pw')
    AppPermission.objects.create(
        user=user, app_name=AppRegistry.objects.get(app_name='quote'), has_access=True,
    )
    return user


def _award(n, **overrides):
    fields = dict(
        sol_number=f'SOL{n}',
        notice_id=f'NOTICE-{n}',
        award_date=date(2026, 1, 15),
        awardee_cage='1ABC2',
        award_basic_number='SPE1-TEST',
        nsn='5340011234567',
        nomenclature='WIDGET',
        source=DibbsAward.SOURCE_DIBBS_FILE,
    )
    fields.update(overrides)
    return DibbsAward.objects.create(**fields)


@override_settings(REQUIRE_LOGIN=True)
class SupplierResearchAccessTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = _quote_user()
        self.no_access = User.objects.create_user('blocked')

    def test_anonymous_denied_page(self):
        response = self.client.get(reverse('quote:supplier_research'))
        self.assertIn(response.status_code, (302, 403))

    def test_anonymous_denied_export(self):
        response = self.client.get(reverse('quote:supplier_research_export', args=['1ABC2']))
        self.assertIn(response.status_code, (302, 403))

    def test_anonymous_denied_panel(self):
        response = self.client.get(reverse('quote:supplier_research_panel', args=['1ABC2', 'status']))
        self.assertIn(response.status_code, (302, 403))

    def test_user_without_quote_access_denied_page(self):
        self.client.force_login(self.no_access)
        response = self.client.get(reverse('quote:supplier_research'))
        self.assertNotEqual(response.status_code, 200)

    def test_user_without_quote_access_denied_export(self):
        self.client.force_login(self.no_access)
        response = self.client.get(reverse('quote:supplier_research_export', args=['1ABC2']))
        self.assertNotEqual(response.status_code, 200)

    def test_user_without_quote_access_denied_panel(self):
        self.client.force_login(self.no_access)
        response = self.client.get(reverse('quote:supplier_research_panel', args=['1ABC2', 'status']))
        self.assertNotEqual(response.status_code, 200)


class CageLookupFormTests(TestCase):
    def test_strip_and_uppercase(self):
        form = CageLookupForm(data={'cage': ' 1abc2 '})
        self.assertTrue(form.is_valid())
        self.assertEqual(form.cleaned_data['cage'], '1ABC2')

    def test_invalid_lengths_and_chars(self):
        for raw in ('1AB', '1AB-2', '123456'):
            form = CageLookupForm(data={'cage': raw})
            self.assertFalse(form.is_valid(), msg=raw)

    def test_widget_is_compact(self):
        html = str(CageLookupForm()['cage'])
        self.assertNotIn('form-control-lg', html)
        self.assertIn('maxlength="5"', html)


SERVICE_NAMES = (
    'find_existing_supplier',
    'get_sam_entity',
    'get_approved_sources',
    'get_award_summary',
    'get_awards',
)


@override_settings(REQUIRE_LOGIN=True)
class SupplierResearchShellTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.client.force_login(_quote_user())

    def test_empty_get_renders_form(self):
        response = self.client.get(reverse('quote:supplier_research'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Supplier Research')
        self.assertContains(response, 'CAGE Code')
        self.assertNotContains(response, 'data-panel-url')

    def test_invalid_cage_shows_error_without_panels(self):
        response = self.client.get(reverse('quote:supplier_research'), {'cage': '1AB-2'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'CAGE must be exactly 5 letters or digits.')
        self.assertNotContains(response, 'data-panel-url')

    def test_valid_cage_renders_four_panels_and_calls_no_service(self):
        patches = [patch(f'quote.services.supplier_research.{name}') for name in SERVICE_NAMES]
        mocks = [p.start() for p in patches]
        try:
            response = self.client.get(reverse('quote:supplier_research'), {'cage': '1abc2'})
        finally:
            for p in patches:
                p.stop()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '1ABC2')
        self.assertContains(response, 'data-panel-url=', count=4)
        for key in PANEL_KEYS:
            self.assertContains(response, f'/quote/research/1ABC2/panel/{key}/')
        self.assertContains(response, 'quote/js/supplier_research.js')
        for mock in mocks:
            mock.assert_not_called()


@override_settings(REQUIRE_LOGIN=True)
class SupplierResearchPanelViewTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.client.force_login(_quote_user())

    def _url(self, panel, cage='1ABC2'):
        return reverse('quote:supplier_research_panel', args=[cage, panel])

    def test_invalid_cage_is_400(self):
        self.assertEqual(self.client.get(self._url('status', cage='1AB-2')).status_code, 400)

    def test_unknown_panel_is_404(self):
        self.assertEqual(self.client.get(self._url('nope')).status_code, 404)

    @patch('quote.services.supplier_research.get_approved_sources')
    @patch('quote.services.supplier_research.get_awards')
    @patch('quote.services.supplier_research.get_award_summary')
    @patch('quote.services.supplier_research.get_sam_entity')
    @patch('quote.services.supplier_research.find_existing_supplier')
    def test_each_panel_renders(self, mock_find, mock_sam, mock_summary, mock_awards, mock_as):
        mock_find.return_value = {
            'id': 1, 'name': 'Acme Defense', 'detail_url': '/suppliers/1/', 'matched_via': 'supplier',
        }
        mock_sam.return_value = SAM_NOT_FOUND
        mock_summary.return_value = EMPTY_SUMMARY
        mock_awards.return_value = []
        mock_as.return_value = {'rows': [], 'has_company_name': False}

        expected = {
            'status': 'Existing STATZ Supplier',
            'sam': 'No SAM.gov record found',
            'awards': 'No award records for this CAGE.',
            'approved': 'No approved-source records for this CAGE.',
        }
        for panel, text in expected.items():
            response = self.client.get(self._url(panel))
            self.assertEqual(response.status_code, 200, msg=panel)
            self.assertContains(response, text, msg_prefix=panel)

    @patch('quote.services.supplier_research.find_existing_supplier', return_value=None)
    def test_status_panel_not_a_supplier(self, _mock):
        response = self.client.get(self._url('status'))
        self.assertContains(response, 'Not a current STATZ supplier')

    @patch('quote.services.supplier_research.get_sam_entity')
    def test_sam_panel_error_state_has_retry(self, mock_sam):
        mock_sam.return_value = {'state': 'error', 'rows': [], 'last_fetched': timezone.now()}
        response = self.client.get(self._url('sam'))
        self.assertContains(response, 'SAM.gov lookup failed')
        self.assertContains(response, 'data-panel-reload')
        self.assertContains(response, 'refresh=1')

    @patch('quote.services.supplier_research.get_sam_entity')
    def test_sam_refresh_passes_force_refresh(self, mock_sam):
        mock_sam.return_value = SAM_NOT_FOUND
        self.client.get(self._url('sam'), {'refresh': '1'})
        mock_sam.assert_called_once_with('1ABC2', force_refresh=True)

    @patch('quote.services.supplier_research.get_sam_entity')
    def test_sam_without_refresh_does_not_force(self, mock_sam):
        mock_sam.return_value = SAM_NOT_FOUND
        self.client.get(self._url('sam'))
        mock_sam.assert_called_once_with('1ABC2', force_refresh=False)

    @patch('quote.services.supplier_research.get_approved_sources')
    @patch('quote.services.supplier_research.get_awards')
    @patch('quote.services.supplier_research.get_award_summary')
    def test_refresh_ignored_outside_sam_panel(self, mock_summary, mock_awards, mock_as):
        mock_summary.return_value = EMPTY_SUMMARY
        mock_awards.return_value = []
        mock_as.return_value = {'rows': [], 'has_company_name': False}
        with patch('dibbs.services.sam_entity.get_or_fetch_cage') as mock_fetch:
            self.client.get(self._url('awards'), {'refresh': '1'})
            self.client.get(self._url('approved'), {'refresh': '1'})
        mock_fetch.assert_not_called()
        mock_summary.assert_called_once_with('1ABC2')
        mock_awards.assert_called_once_with('1ABC2', limit=sr.AWARD_DISPLAY_LIMIT)

    @patch('quote.services.supplier_research.get_approved_sources')
    def test_approved_panel_collapses_after_25_rows(self, mock_as):
        rows = [{'nsn': f'534001123{n:04d}', 'part_number': f'PN{n}', 'company_name': None} for n in range(30)]
        mock_as.return_value = {'rows': rows, 'has_company_name': False}
        response = self.client.get(self._url('approved'))
        self.assertContains(response, 'data-extra-row', count=5)
        self.assertContains(response, 'Show all 30')
        self.assertNotContains(response, 'Company Name')

    @patch('quote.services.supplier_research.get_approved_sources')
    def test_approved_panel_short_table_has_no_toggle(self, mock_as):
        mock_as.return_value = {
            'rows': [{'nsn': '5340011234567', 'part_number': 'PN1', 'company_name': 'ACME'}],
            'has_company_name': True,
        }
        response = self.client.get(self._url('approved'))
        self.assertNotContains(response, 'data-show-all')
        self.assertContains(response, 'Company Name')

    @patch('quote.services.supplier_research.get_awards')
    @patch('quote.services.supplier_research.get_award_summary')
    def test_awards_panel_shows_source_and_freshness(self, mock_summary, mock_awards):
        mock_summary.return_value = dict(
            EMPTY_SUMMARY, total_awards=1, distinct_nsns=1, total_value=Decimal('12.50'),
            data_current_through=date(2026, 9, 25),
        )
        mock_awards.return_value = [{
            'award_date': date(2026, 1, 15), 'posted_date': None, 'award_basic_number': 'SPE1-TEST',
            'delivery_order_number': '', 'nsn': '5340011234567', 'nomenclature': 'WIDGET',
            'total_contract_price': Decimal('12.50'), 'source': 'DIBBS_FILE',
        }]
        response = self.client.get(self._url('awards'))
        self.assertContains(response, 'DIBBS_FILE')
        self.assertContains(response, 'Total Value (where reported)')
        self.assertContains(response, 'Award data current through 2026-09-25')


class GetSamEntityTests(TestCase):
    def _record(self, **kw):
        base = dict(
            fetch_error=False, raw_json={}, last_fetched=timezone.now(), entity_name='ACME INC',
            cage_code='1ABC2', website='', physical_address_line1='1 Main', physical_address_line2='',
            physical_city='Town', physical_state='VA', physical_zip='22000', mailing_address='',
            sba_flags=[], naics_codes=[], psc_codes=[],
        )
        base.update(kw)
        return SimpleNamespace(**base)

    def _state(self, record, **kw):
        with patch('dibbs.services.sam_entity.get_or_fetch_cage', return_value=record) as mock:
            result = sr.get_sam_entity('1ABC2', **kw)
        return result, mock

    def test_fetch_error_is_error(self):
        record = self._record(fetch_error=True, raw_json={'error': 'SAM API call failed', 'found': False})
        result, _ = self._state(record)
        self.assertEqual(result['state'], 'error')
        self.assertEqual(result['rows'], [])
        self.assertEqual(result['last_fetched'], record.last_fetched)

    def test_lookup_exception_is_error(self):
        with patch('dibbs.services.sam_entity.get_or_fetch_cage', side_effect=RuntimeError('boom')):
            result = sr.get_sam_entity('1ABC2')
        self.assertEqual(result['state'], 'error')

    def test_found_false_is_not_found(self):
        record = self._record(raw_json={'found': False, 'cage_code': '1ABC2'})
        result, _ = self._state(record)
        self.assertEqual(result['state'], 'not_found')

    def test_populated_is_ok_and_formats_code_lists(self):
        record = self._record(
            raw_json={'found': True, 'uei': 'UEI123456789', 'registration_status': 'Active'},
            naics_codes=[{'code': '332710', 'desc': 'Machine Shops', 'primary': True}],
            psc_codes=[{'code': '5340', 'desc': 'Hardware'}],
        )
        result, _ = self._state(record)
        self.assertEqual(result['state'], 'ok')
        rows = dict(result['rows'])
        self.assertEqual(rows['Legal name'], 'ACME INC')
        self.assertEqual(rows['NAICS codes'], '332710 (Machine Shops)')
        self.assertEqual(rows['PSC codes'], '5340 (Hardware)')

    def test_force_refresh_is_forwarded(self):
        _, mock = self._state(self._record(raw_json={'found': False}), force_refresh=True)
        mock.assert_called_once_with('1ABC2', force_refresh=True)


class SupplierResearchServiceDataTests(TestCase):
    def test_identical_award_rows_collapse(self):
        _award(1)
        _award(2)  # differs only by notice_id / sol_number, which are not dedupe fields
        self.assertEqual(sr.get_award_summary('1ABC2')['total_awards'], 1)
        self.assertEqual(len(sr.get_awards('1ABC2', limit=None)), 1)

    def test_different_awards_do_not_collapse(self):
        _award(1)
        _award(2, award_basic_number='SPE1-OTHER')
        self.assertEqual(len(sr.get_awards('1ABC2', limit=None)), 2)

    def test_total_value_not_double_counted(self):
        _award(1, total_contract_price=Decimal('10.00'))
        _award(2, total_contract_price=Decimal('10.00'))
        _award(3, award_basic_number='SPE1-OTHER', total_contract_price=Decimal('5.00'))
        self.assertEqual(sr.get_award_summary('1ABC2')['total_value'], Decimal('15.00'))

    def test_total_value_none_when_nothing_reported(self):
        _award(1)
        self.assertIsNone(sr.get_award_summary('1ABC2')['total_value'])

    def test_award_limit_applies(self):
        for n in range(3):
            _award(n, award_basic_number=f'SPE1-{n}')
        self.assertEqual(len(sr.get_awards('1ABC2', limit=2)), 2)

    def test_freshness_ignores_future_and_award_date(self):
        today = timezone.localdate()
        _award(1, posted_date=today + timedelta(days=30))
        _award(2, award_basic_number='B', posted_date=today - timedelta(days=1),
               award_date=today + timedelta(days=400))
        summary = sr.get_award_summary('1ABC2')
        self.assertEqual(summary['data_current_through'], today - timedelta(days=1))

    def test_faux_rows_excluded_everywhere(self):
        today = timezone.localdate()
        _award(1)
        _award(2, award_basic_number='FAUX', is_faux=True, award_date=date(2099, 9, 30),
               posted_date=today - timedelta(days=2))
        summary = sr.get_award_summary('1ABC2')
        self.assertEqual(summary['total_awards'], 1)
        self.assertEqual(summary['last_award'], date(2026, 1, 15))
        self.assertEqual(len(sr.get_awards('1ABC2', limit=None)), 1)
        # a faux row's posted date must not drive freshness either
        self.assertIsNone(summary['data_current_through'])

    def test_awards_scoped_to_cage(self):
        _award(1)
        _award(2, awardee_cage='9ZZZ9')
        self.assertEqual(len(sr.get_awards('1ABC2', limit=None)), 1)

    def test_identical_approved_rows_collapse(self):
        for _ in range(2):
            ApprovedSource.objects.create(nsn='5340011234567', approved_cage='1ABC2', part_number='PN1')
        ApprovedSource.objects.create(nsn='5340011234567', approved_cage='9ZZZ9', part_number='PN1')
        result = sr.get_approved_sources('1ABC2')
        self.assertEqual(len(result['rows']), 1)
        self.assertFalse(result['has_company_name'])

    def test_approved_has_company_name(self):
        ApprovedSource.objects.create(
            nsn='5340011234567', approved_cage='1ABC2', part_number='PN1', company_name='ACME',
        )
        self.assertTrue(sr.get_approved_sources('1ABC2')['has_company_name'])


@override_settings(REQUIRE_LOGIN=True)
class SupplierResearchExportTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.client.force_login(_quote_user())

    def _export(self, sam=SAM_NOT_FOUND, approved=None):
        approved = approved or {'rows': [], 'has_company_name': False}
        awards = [{
            'award_date': None, 'posted_date': None, 'award_basic_number': 'SPE1-TEST',
            'delivery_order_number': '', 'nsn': '5340011234567', 'nomenclature': 'WIDGET',
            'total_contract_price': None, 'source': 'DIBBS_FILE',
        }]
        with patch('quote.services.supplier_research.find_existing_supplier', return_value=None), \
                patch('quote.services.supplier_research.get_sam_entity', return_value=sam), \
                patch('quote.services.supplier_research.get_approved_sources', return_value=approved), \
                patch('quote.services.supplier_research.get_award_summary',
                      return_value=dict(EMPTY_SUMMARY, total_awards=1, distinct_nsns=1)), \
                patch('quote.services.supplier_research.get_awards', return_value=awards):
            return self.client.get(reverse('quote:supplier_research_export', args=['1ABC2']))

    def test_export_workbook(self):
        response = self._export()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('Supplier_Research_1ABC2_', response['Content-Disposition'])
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Supplier Data', 'Approved Sources', 'Awards'])
        awards_ws = wb['Awards']
        headers = [c.value for c in awards_ws[1]]
        self.assertIn('Source', headers)
        nsn_cell = awards_ws.cell(row=2, column=headers.index('NSN') + 1)
        self.assertEqual(nsn_cell.value, '5340011234567')
        self.assertIsInstance(nsn_cell.value, str)
        self.assertEqual(awards_ws.cell(row=2, column=headers.index('Source') + 1).value, 'DIBBS_FILE')

    def test_approved_sources_sheet_has_no_cage_column(self):
        approved = {
            'rows': [{'nsn': '5340011234567', 'part_number': 'PN1', 'company_name': None}],
            'has_company_name': False,
        }
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(self._export(approved=approved).content))
        headers = [c.value for c in wb['Approved Sources'][1]]
        self.assertEqual(headers, ['NSN', 'Part number'])
        self.assertEqual(wb['Approved Sources'].cell(row=2, column=2).value, 'PN1')

    def test_sam_error_and_not_found_write_explanatory_rows(self):
        from openpyxl import load_workbook

        def sam_text(sam):
            wb = load_workbook(BytesIO(self._export(sam=sam).content))
            return [r[1].value for r in wb['Supplier Data'].iter_rows() if r[0].value == 'SAM.gov']

        self.assertEqual(sam_text(SAM_NOT_FOUND), ['No SAM.gov record found'])
        error = {'state': 'error', 'rows': [], 'last_fetched': timezone.now()}
        self.assertIn('lookup failed', sam_text(error)[0])

    def test_invalid_cage_export_400(self):
        response = self.client.get(reverse('quote:supplier_research_export', args=['NOPE']))
        self.assertEqual(response.status_code, 400)
